"""
Tender analysis — what a tender asks for, for the opportunity pages.

  1. Award criteria (no LLM): type and weight straight from the notice's
     eForms XML, e.g. price 30 % / quality 50 % / quality 20 %.
  2. Tender Analysis agent: reads the notice and the text of the tender
     documents the buyer published (tender conditions, specifications,
     proof forms, Q&A catalogues) and writes an English summary of the
     tender plus every requirement, grouped as eligibility, project
     references, human resources and technical / project requirements.

Runs in the same pipeline run that adds an opportunity, so its page shows the
summary and requirements together; re-runs when new clarifications, Q&A or
specifications are published. It describes the tender only — WalliD's fit is
the Tender Evaluation Agent's job (webapp/lib/tenderEvaluation.ts), run from
the opportunity page.
"""
import io
import json
import os
import re
import sys
import uuid
import zipfile
from datetime import datetime

from sqlalchemy import text

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from services.llm_client import call_llm_json_premium, load_prompt  # noqa: E402

_ROOT = os.path.join(os.path.dirname(__file__), "..")
CATEGORIES = {"LEGAL", "FINANCIAL", "TURNOVER", "INSURANCE", "CERTIFICATION", "COMPANY_EXPERIENCE", "REFERENCE",
              "TEAM", "CV", "EDUCATION", "PERSONAL_CERTIFICATION", "SECURITY_CLEARANCE", "LANGUAGE", "FTE",
              "LOCAL_PRESENCE", "TECHNICAL", "SECURITY", "PRIVACY", "EIDAS", "EUDI", "INTEROPERABILITY", "HOSTING",
              "SLA", "IMPLEMENTATION", "CONSORTIUM", "SUBCONTRACTING", "EVIDENCE"}


def _clean(s):
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", s or "")).strip()


def notice_criteria(xml: str) -> dict:
    """Selection criteria, tenderer requirements and award criteria from eForms XML."""
    selection = [_clean(d) for b in re.findall(r"<efac:SelectionCriteria>(.*?)</efac:SelectionCriteria>", xml, re.S)
                 for d in re.findall(r"<cbc:Description[^>]*>(.*?)</cbc:Description>", b, re.S)]
    tenderer = [_clean(d) for b in re.findall(r"<cac:SpecificTendererRequirement>(.*?)</cac:SpecificTendererRequirement>", xml, re.S)
                for d in re.findall(r"<cbc:Description[^>]*>(.*?)</cbc:Description>", b, re.S)]
    award = []
    for b in re.findall(r"<cac:SubordinateAwardingCriterion>(.*?)</cac:SubordinateAwardingCriterion>", xml, re.S):
        kind = re.search(r"<cbc:AwardingCriterionTypeCode[^>]*>([^<]+)<", b)
        weight = re.search(r"<efbc:ParameterNumeric>([\d.]+)<", b)
        name = re.search(r"<cbc:Name[^>]*>(.*?)</cbc:Name>", b, re.S)
        desc = re.search(r"<cbc:Description[^>]*>(.*?)</cbc:Description>", b, re.S)
        award.append({"type": kind.group(1) if kind else None, "weight": float(weight.group(1)) if weight else None,
                      "name": _clean(name.group(1)) if name else None, "description": _clean(desc.group(1)) if desc else None})
    dedupe = lambda xs: [x for x in dict.fromkeys(xs) if len(x) > 15]
    return {"selection": dedupe(selection), "tenderer": dedupe(tenderer), "award": award}


def store_award_criteria(session, opportunity_id: str, award: list) -> int:
    """Deterministic part: replace the opportunity's award criteria with the notice's."""
    if not award:
        return 0
    session.execute(text("delete from award_criteria where opportunity_id = :o"), {"o": opportunity_id})
    for a in award:
        label = {"price": "Price", "cost": "Cost", "quality": "Quality"}.get(a["type"] or "", (a["type"] or "Criterion").title())
        session.execute(text("insert into award_criteria (opportunity_id, criterion, weight, subcriteria) values (:o, :c, :w, cast(:s as jsonb))"),
                        {"o": opportunity_id, "c": label, "w": a["weight"],
                         "s": json.dumps({"type": a["type"], "name": a["name"], "description": a["description"]})})
    session.commit()
    return len(award)


GROUPS = {"ELIGIBILITY", "REFERENCES", "HUMAN_RESOURCES", "TECHNICAL"}
# Documents worth reading, in order; contracts and price sheets add little.
READ_ORDER = ["TENDER_SPECIFICATIONS", "TECHNICAL_SPECIFICATIONS", "AWARD_CRITERIA", "FORM", "Q_AND_A", "CLARIFICATION", "ANNEX"]
DOC_CHARS, TOTAL_CHARS = 60_000, 260_000
MAIN_DOC_CHARS = 100_000  # tender conditions and specifications get more room


def notice_description(xml: str) -> str:
    """The procedure's and each lot's name, description and duration."""
    out = []
    for b in re.findall(r"<cac:ProcurementProject>(.*?)</cac:ProcurementProject>", xml, re.S):
        name = re.search(r"<cbc:Name[^>]*>(.*?)</cbc:Name>", b, re.S)
        desc = re.search(r"<cbc:Description[^>]*>(.*?)</cbc:Description>", b, re.S)
        dur = re.search(r'<cbc:DurationMeasure unitCode="(\w+)">(\d+)<', b)
        out.append(" — ".join(x for x in [_clean(name.group(1)) if name else "", _clean(desc.group(1)) if desc else "",
                                         f"duration {dur.group(2)} {dur.group(1).lower()}" if dur else ""] if x))
    return "\n".join(dict.fromkeys(x for x in out if x))


def _file_text(name: str, data: bytes) -> str:
    """Plain text of a PDF, Word or Excel file (or of those inside a zip)."""
    low = name.lower()
    try:
        if data[:4] == b"%PDF" or low.endswith(".pdf"):
            from pypdf import PdfReader
            return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages[:80])
        if data[:2] == b"PK":
            z = zipfile.ZipFile(io.BytesIO(data))
            names = z.namelist()
            if "word/document.xml" in names:
                from docx import Document as Docx
                d = Docx(io.BytesIO(data))
                rows = [" | ".join(c.text.strip() for c in r.cells) for t in d.tables for r in t.rows]
                return "\n".join([p.text for p in d.paragraphs if p.text.strip()] + rows)
            if "xl/workbook.xml" in names:
                from openpyxl import load_workbook
                wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
                return "\n".join(" | ".join(str(v) for v in row if v is not None)
                                 for ws in wb.worksheets for row in ws.iter_rows(values_only=True) if any(row))
            inner = [n for n in names if n.lower().endswith((".pdf", ".docx", ".xlsx"))][:8]
            return "\n\n".join(f"[{n}]\n{_file_text(n, z.read(n))}" for n in inner)
    except Exception:
        return ""
    return ""


def _recency(name: str) -> int:
    """Sort key from a file name's date (Stand 01.10.2026, 2026-10-01) or number (Clarification Note 31)."""
    d = re.search(r"(\d{1,2})\.(\d{1,2})\.(20\d\d)", name) or re.search(r"(20\d\d)-(\d\d)-(\d\d)", name)
    if d:
        a, b, c = (int(x) for x in d.groups())
        return (c * 10000 + b * 100 + a) if c > 1000 else (a * 10000 + b * 100 + c)
    n = re.findall(r"\d+", name)
    return int(n[-1]) if n else 0


def document_texts(session, opportunity_id: str) -> list:
    """(name, text) of the downloadable tender documents, most useful first, within budget."""
    import requests
    rows = session.execute(text("""select name, document_type, url from documents where opportunity_id = :o
            and (url like '%downloadContractDocument%' or url like '%/de/documents/%')"""), {"o": opportunity_id}).fetchall()
    rank = {t: i for i, t in enumerate(READ_ORDER)}
    rows = sorted((r for r in rows if r.document_type in rank), key=lambda r: (rank[r.document_type], -_recency(r.name)))
    picked, seen_kind, total = [], {}, 0
    for r in rows:
        kind = r.document_type
        if kind in ("Q_AND_A", "CLARIFICATION") and seen_kind.get(kind, 0) >= 4:
            continue  # catalogues are cumulative: the latest few say it all
        try:
            resp = requests.get(r.url, timeout=60, headers={"User-Agent": "Mozilla/5.0 TenderTown/1.0"})
            if resp.status_code != 200 or len(resp.content) > 30_000_000:
                continue
        except Exception:
            continue
        body = re.sub(r"[ \t]+", " ", _file_text(r.name, resp.content)).strip()[:MAIN_DOC_CHARS if kind in READ_ORDER[:2] else DOC_CHARS]
        if len(body) < 200:
            continue
        if total + len(body) > TOTAL_CHARS:
            body = body[:max(0, TOTAL_CHARS - total)]
        if not body:
            break
        picked.append((r.name, body))
        seen_kind[kind] = seen_kind.get(kind, 0) + 1
        total += len(body)
    return picked


def extract_requirements(description: str, criteria: dict, doc_names: list, docs: list) -> dict:
    listing = "\n".join([
        "## Notice", "Procurement description:", description or "(none)",
        "Selection criteria:", *[f"- {s[:2000]}" for s in criteria["selection"]],
        "Other tenderer requirements:", *[f"- {s[:2000]}" for s in criteria["tenderer"]],
        "Award criteria:", *[f"- {a['type']} {a['weight']}%: {a['name'] or ''} {a['description'] or ''}" for a in criteria["award"]],
        "All published documents:", *[f"- {d}" for d in doc_names[:80]],
        *[f"\n## Document: {name}\n{body}" for name, body in docs],
    ])
    return call_llm_json_premium(load_prompt("tender_requirements.md"), listing, max_tokens=16000, effort="medium")


def page_text(url: str) -> str:
    """Readable text of a non-TED opportunity's own page (Tavily's extractor when configured)."""
    import requests
    key = os.environ.get("TAVILY_API_KEY")
    if key:
        try:
            r = requests.post("https://api.tavily.com/extract", json={"urls": [url]},
                              headers={"Authorization": f"Bearer {key}"}, timeout=40)
            body = ((r.json().get("results") or [{}])[0] or {}).get("raw_content") or ""
            if len(body) > 300:
                return body[:DOC_CHARS]
        except Exception:
            pass
    try:
        page = requests.get(url, timeout=40, headers={"User-Agent": "Mozilla/5.0 TenderTown/1.0"}).text
    except Exception:
        return ""
    page = re.sub(r"<(script|style|noscript)[^>]*>.*?</\1>", " ", page, flags=re.S | re.I)
    return _clean(page)[:DOC_CHARS]


def analyse_tenders(session, errors: list, limit: int = 40) -> dict:
    """Award criteria for every tracked TED notice; summary + requirements (LLM)
    for every active opportunity not analysed yet, or with new clarifications /
    Q&A / specifications since its last analysis — so a page never has a
    summary without its requirements."""
    import requests
    opps = session.execute(text("""select opportunity_id, official_url, status, deadline, tender_analysed_at,
            exists (select 1 from documents d where d.opportunity_id = o.opportunity_id and d.downloaded_at > o.tender_analysed_at
                    and d.document_type in ('CLARIFICATION','Q_AND_A','CORRIGENDUM','TENDER_SPECIFICATIONS','TECHNICAL_SPECIFICATIONS')) as new_docs
            from opportunities o where official_url is not null""")).fetchall()
    done = {"award": 0, "requirements": 0}
    llm_budget = limit
    for o in opps:
        ted = "ted.europa.eu" in o.official_url
        xml, crit = "", {"selection": [], "tenderer": [], "award": []}
        active = o.status in ("OPEN", "SIGNAL", "UNVERIFIED") and (o.deadline is None or o.deadline >= datetime.utcnow())
        if ted:
            pub = o.official_url.rstrip("/").split("/")[-1]
            try:
                xml = requests.get(f"https://ted.europa.eu/en/notice/{pub}/xml", timeout=40).text
            except Exception as e:
                errors.append(f"tender analysis {o.opportunity_id}: {e}"[:200]); continue
            crit = notice_criteria(xml)
            done["award"] += store_award_criteria(session, o.opportunity_id, crit["award"])
        if not active or (o.tender_analysed_at and not o.new_docs) or llm_budget <= 0:
            continue
        llm_budget -= 1
        names = [r.name for r in session.execute(text("select name from documents where opportunity_id = :o"), {"o": o.opportunity_id})]
        docs = document_texts(session, o.opportunity_id)
        description = notice_description(xml) if ted else ""
        if not ted:
            body = page_text(o.official_url)
            if body:
                docs = [("Opportunity page " + o.official_url, body)] + docs
        try:
            out = extract_requirements(description, crit, names, docs)
        except Exception as e:
            errors.append(f"requirements {o.opportunity_id}: {e}"[:300])
            if "credit" in str(e).lower():
                break
            continue
        session.execute(text("delete from requirement_matches where requirement_id in (select requirement_id from requirements where opportunity_id = :o)"), {"o": o.opportunity_id})
        session.execute(text("delete from requirements where opportunity_id = :o"), {"o": o.opportunity_id})
        for r in out.get("requirements") or []:
            cat, group = (r.get("category") or "").upper(), (r.get("group") or "").upper()
            if not r.get("text") or group not in GROUPS:
                continue
            session.execute(text("""insert into requirements (requirement_id, opportunity_id, requirement_group, category, requirement_text,
                    mandatory, threshold, evidence_required, document, source_url, confidence)
                    values (:id, :o, :g, :c, :t, :m, :th, :ev, :doc, :u, 'INFERRED')"""),
                            {"id": f"{o.opportunity_id}-{uuid.uuid4().hex[:8]}", "o": o.opportunity_id, "g": group,
                             "c": cat if cat in CATEGORIES else None, "t": r["text"][:2000], "m": r.get("mandatory") is not False,
                             "th": (r.get("threshold") or None) and str(r["threshold"])[:256], "ev": r.get("evidence") or None,
                             "doc": (r.get("source") or None) and str(r["source"])[:512], "u": o.official_url})
        for a in out.get("award_criteria") or []:
            if a.get("criterion"):
                session.execute(text("""update award_criteria set subcriteria = subcriteria || cast(:s as jsonb)
                        where opportunity_id = :o and weight is not distinct from :w"""),
                                {"s": json.dumps({"name_en": a["criterion"], "description_en": a.get("description")}),
                                 "o": o.opportunity_id, "w": float(a["weight"]) if a.get("weight") is not None else None})
        session.execute(text("update opportunities set tender_summary = :s, tender_analysed_at = now() where opportunity_id = :o"),
                        {"s": (out.get("summary") or "").strip() or None, "o": o.opportunity_id})
        session.commit()
        done["requirements"] += 1
    return done
